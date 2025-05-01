import tkinter as tk 
from tkinter import messagebox
from openpyxl import load_workbook, Workbook

window = tk. Tk()
window.title("User Score Entry")
window.geometry('300x250')
window.configure(bg="azure3")

def cre_excel():
    wb = Workbook()
    ws = wb.active
    ws.title = "userdata"
    ws.append(["Name", "Score","Remarks"])

def edit_excel():
    if not bug():
        return
    nm = name_entry.get().strip()
    sc = int(score_entry.get())
    
    try:
        wb = load_workbook("data.xlsx")
        if "userdata" in wb.sheetnames:
            ws = wb["userdata"]
        else:
            ws = wb.active
    except FileNotFoundError:
        wb = Workbook()
        ws = wb.active
        ws.title = "userdata"
        ws.append(["Name", "Score","Remarks"])

    ws = wb.active
    if sc >= 75 and sc <= 100:
        ws.append([nm, sc, "Passed"])
        messagebox.showinfo(title="Success", message="Exellent Work")
    elif sc >= 0 and sc <= 74:
        ws.append([nm, sc, "Failed"])
        messagebox.showinfo(title="Success", message="Study Well Next Time")   
    else:
        messagebox.showerror(title="Error", message="Invalid Input, Please Enter your score between 0 to 100")

    wb.save("data.xlsx")
    messagebox.showinfo(title="Success", message="Saved Successfully")

def up_excel():
    if not bug():
        return
    nm = name_entry.get().strip()
    sc = int(score_entry.get())

    wb = load_workbook("data.xlsx")
    ws = wb["userdata"]

    for row in ws.iter_rows(min_row=2):
        row_val = row[0].value
        if row_val == nm:
            row[1].value = sc
            if sc >= 75 and sc <= 100:
                row[2].value = "Passed"
            elif sc >= 0 and sc <= 74:
                row[2].value = "Failed"
            wb.save("data.xlsx")
            messagebox.showinfo(title="Success", message="Updated Successfully")
            break

def bug():
    
    try:
        sc = int(score_entry.get())
        if sc >= 75 and sc <= 100:
            messagebox.showinfo(title="Success", message="Exellent Work")
        elif sc >= 0 and sc <= 74:
            messagebox.showinfo(title="Success", message="Study Well Next Time")   
        else:
            messagebox.showerror(title="Error", message="Invalid Input, Please Enter your score between 0 to 100")
    except ValueError:
        messagebox.showerror(title="Error", message="Invalid Input, Please Enter your score between 0 to 100")
        return False


frm = tk. Frame(window, bg="azure3")
frm.place(relx = 0.5, rely = 0.5, anchor = "center")


tk. Label (frm, text="Name", bg="azure3", font=("Arial", 10, "bold")).grid(row=1, column=0)
tk. Label(frm, text="Score", bg="azure3", font=("Arial", 10, "bold")).grid(row=3, column=0)


name_entry = tk. Entry(frm)
score_entry= tk.Entry (frm)

name_entry.grid(row=0, column=0)
score_entry.grid(row= 2, column=0)

sub = tk.Button(frm, text="Submit",command=edit_excel)
up = tk.Button(frm, text="Update", command=up_excel)
sub.grid(row=6, column=0, sticky="w")
up.grid(row=6, column=0, sticky="e")



frm.mainloop()